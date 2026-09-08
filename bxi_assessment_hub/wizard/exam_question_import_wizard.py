# -*- coding: utf-8 -*-
# Copyright (c) 2026 BXI Technology Pvt. Ltd. All Rights Reserved.
# License OPL-1 (Odoo Proprietary License v1.0, see LICENSE file for full text).
import base64
import io
import re

from odoo import _, fields, models
from odoo.exceptions import UserError

QUESTION_TYPE_ALIASES = {
    'mcq': 'mcq',
    'multiple choice': 'mcq',
    'short answer': 'short_answer',
    'short_answer': 'short_answer',
    'short': 'short_answer',
    'long answer': 'long_answer',
    'long_answer': 'long_answer',
    'long': 'long_answer',
}
OPTION_COLUMNS = ['Option A', 'Option B', 'Option C', 'Option D']


def _normalize_type(raw):
    key = (raw or '').strip().lower()
    question_type = QUESTION_TYPE_ALIASES.get(key)
    if not question_type:
        raise UserError(_('Unknown question type "%s". Use MCQ, Short Answer or Long Answer.') % raw)
    return question_type


class BxiExamQuestionImportWizard(models.TransientModel):
    _name = 'bxi.exam.question.import.wizard'
    _description = 'Import Exam Questions'

    exam_id = fields.Many2one('bxi.exam', string='Exam', required=True)
    upload_file = fields.Binary(string='File', required=True)
    file_name = fields.Char(string='File Name')

    def action_download_template(self):
        return {
            'type': 'ir.actions.act_url',
            'url': '/bxi_assessment_hub/static/xlsx/question_import_template.xlsx',
            'target': 'new',
        }

    def action_import(self):
        self.ensure_one()
        if not self.file_name:
            raise UserError(_('Choose a file to import.'))
        extension = self.file_name.rsplit('.', 1)[-1].lower()
        content = base64.b64decode(self.upload_file)

        if extension == 'xlsx':
            rows = self._parse_xlsx(content)
        elif extension == 'docx':
            rows = self._parse_docx(content)
        elif extension == 'pdf':
            rows = self._parse_pdf(content)
        else:
            raise UserError(_(
                'Unsupported file type ".%s". Upload an Excel (.xlsx), Word (.docx) or PDF (.pdf) file.'
            ) % extension)

        if not rows:
            raise UserError(_('No questions were found in this file.'))

        Question = self.env['bxi.exam.question']
        Option = self.env['bxi.exam.question.option']
        for row in rows:
            question = Question.create({
                'exam_id': self.exam_id.id,
                'question_type': row['question_type'],
                'question_text': row['question_text'],
                'marks': row['marks'],
                'sample_answer_keywords': row.get('sample_answer_keywords') or False,
            })
            for option in row.get('options') or []:
                Option.create({
                    'question_id': question.id,
                    'option_text': option['option_text'],
                    'is_correct': option['is_correct'],
                })

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Questions Imported'),
                'message': _('%d question(s) added to "%s".') % (len(rows), self.exam_id.name),
                'type': 'success',
                'next': {
                    'type': 'ir.actions.act_window',
                    'res_model': 'bxi.exam',
                    'res_id': self.exam_id.id,
                    'views': [[False, 'form']],
                },
            },
        }

    # --- Excel: the recommended, fully structured import format ---

    def _parse_xlsx(self, content):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = [str(cell).strip() if cell else '' for cell in next(rows_iter)]
        except StopIteration:
            return []
        col_index = {name: idx for idx, name in enumerate(header)}

        required = ['Question Text', 'Type', 'Marks']
        missing = [col for col in required if col not in col_index]
        if missing:
            raise UserError(_('The Excel file is missing required column(s): %s') % ', '.join(missing))

        def get(row, name):
            idx = col_index.get(name)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        rows = []
        for raw_row in rows_iter:
            if raw_row is None or not any(raw_row):
                continue
            question_text = get(raw_row, 'Question Text')
            if not question_text:
                continue
            question_type = _normalize_type(get(raw_row, 'Type'))
            marks = get(raw_row, 'Marks') or 0
            entry = {
                'question_type': question_type,
                'question_text': str(question_text).strip(),
                'marks': float(marks),
                'sample_answer_keywords': get(raw_row, 'Sample Answer/Keywords'),
            }
            if question_type == 'mcq':
                option_texts = [get(raw_row, col) for col in OPTION_COLUMNS]
                correct_label = str(get(raw_row, 'Correct Option') or '').strip().upper()
                options = []
                for idx, text in enumerate(option_texts):
                    if not text:
                        continue
                    label = chr(65 + idx)
                    options.append({'option_text': str(text).strip(), 'is_correct': label == correct_label})
                if len(options) < 2:
                    raise UserError(_('Question "%s" needs at least two answer options.') % entry['question_text'])
                if len([o for o in options if o['is_correct']]) != 1:
                    raise UserError(_(
                        'Question "%s" must have exactly one correct option (see the Correct Option column).'
                    ) % entry['question_text'])
                entry['options'] = options
            rows.append(entry)
        return rows

    # --- Word / PDF: both reduce to plain text, parsed with the same
    # lightweight block format (documented on the wizard form) rather than
    # guessed at, since freeform prose can't be reliably turned into
    # structured questions:
    #
    #   Q: <question text>
    #   Type: mcq | short_answer | long_answer
    #   Marks: <number>
    #   Answer: <sample answer / keywords>        (short_answer/long_answer)
    #   A) <option text>                           (mcq, one line per option)
    #   B) <option text>
    #   Correct: A                                 (mcq)
    #
    # Blocks are separated by one or more blank lines.

    def _parse_docx(self, content):
        import docx

        document = docx.Document(io.BytesIO(content))
        text = '\n'.join(p.text for p in document.paragraphs)
        return self._parse_text_blocks(text)

    def _parse_pdf(self, content):
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(content))
        text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        return self._parse_text_blocks(text)

    def _parse_text_blocks(self, text):
        blocks = re.split(r'\n\s*\n', text.strip())
        rows = []
        for block in blocks:
            lines = [line.strip() for line in block.splitlines() if line.strip()]
            if not lines:
                continue
            entry = {'question_type': 'short_answer', 'marks': 1.0, 'options': []}
            question_lines = []
            for line in lines:
                match_option = re.match(r'^([A-Da-d])[\).]\s*(.+)$', line)
                if re.match(r'^Q\s*[:\-]', line, re.IGNORECASE):
                    question_lines.append(re.sub(r'^Q\s*[:\-]\s*', '', line, flags=re.IGNORECASE))
                elif re.match(r'^Type\s*[:\-]', line, re.IGNORECASE):
                    entry['question_type'] = _normalize_type(
                        re.sub(r'^Type\s*[:\-]\s*', '', line, flags=re.IGNORECASE))
                elif re.match(r'^Marks\s*[:\-]', line, re.IGNORECASE):
                    raw_marks = re.sub(r'^Marks\s*[:\-]\s*', '', line, flags=re.IGNORECASE)
                    try:
                        entry['marks'] = float(raw_marks)
                    except ValueError:
                        pass
                elif re.match(r'^Answer\s*[:\-]', line, re.IGNORECASE):
                    entry['sample_answer_keywords'] = re.sub(r'^Answer\s*[:\-]\s*', '', line, flags=re.IGNORECASE)
                elif re.match(r'^Correct\s*[:\-]', line, re.IGNORECASE):
                    entry['_correct_label'] = re.sub(
                        r'^Correct\s*[:\-]\s*', '', line, flags=re.IGNORECASE).strip().upper()
                elif match_option:
                    entry['options'].append({
                        'label': match_option.group(1).upper(),
                        'option_text': match_option.group(2).strip(),
                    })
                else:
                    question_lines.append(line)
            if not question_lines:
                continue
            entry['question_text'] = ' '.join(question_lines).strip()
            if entry['question_type'] == 'mcq':
                correct_label = entry.pop('_correct_label', None)
                for option in entry['options']:
                    option['is_correct'] = option.pop('label') == correct_label
                if len(entry['options']) < 2:
                    raise UserError(_(
                        'Question "%s" needs at least two answer options (lines starting "A)", "B)", ...).'
                    ) % entry['question_text'])
                if len([o for o in entry['options'] if o['is_correct']]) != 1:
                    raise UserError(_(
                        'Question "%s" must have exactly one correct option ("Correct: A").'
                    ) % entry['question_text'])
            else:
                entry.pop('options', None)
                entry.pop('_correct_label', None)
            rows.append(entry)
        return rows
