# -*- coding: utf-8 -*-
# Copyright (c) 2026 BXI Technology Pvt. Ltd. All Rights Reserved.
# License OPL-1 (Odoo Proprietary License v1.0, see LICENSE file for full text).
import base64
import io

from odoo.exceptions import UserError
from odoo.tests.common import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestExamQuestionImportWizard(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.course = cls.env['op.course'].create({'name': 'Class 9', 'code': 'C9'})
        cls.subject = cls.env['op.subject'].create({'name': 'Chemistry', 'code': 'CHEM'})
        cls.teacher = cls.env['op.faculty'].create({
            'first_name': 'Anita', 'last_name': 'Rao',
            'birth_date': '1988-02-02', 'gender': 'female',
        })
        cls.exam = cls.env['bxi.exam'].create({
            'name': 'Chemistry Practical Assessment',
            'class_id': cls.course.id,
            'subject_id': cls.subject.id,
            'teacher_id': cls.teacher.id,
        })

    def _wizard_for(self, content, file_name):
        return self.env['bxi.exam.question.import.wizard'].create({
            'exam_id': self.exam.id,
            'upload_file': base64.b64encode(content),
            'file_name': file_name,
        })

    # --- Excel ---

    def _build_xlsx(self, rows, headers=None):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(headers or [
            'Question Text', 'Type', 'Marks', 'Option A', 'Option B', 'Option C', 'Option D',
            'Correct Option', 'Sample Answer/Keywords',
        ])
        for row in rows:
            ws.append(row)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def test_xlsx_import_creates_short_answer_and_mcq(self):
        content = self._build_xlsx([
            ["State Newton's Third Law", 'Short Answer', 5, '', '', '', '', '', 'action, reaction'],
            ['Who formulated the laws of motion?', 'MCQ', 1, 'Newton', 'Jonas', 'Frost', '', 'A', ''],
        ])
        wizard = self._wizard_for(content, 'questions.xlsx')
        wizard.action_import()

        self.assertEqual(len(self.exam.question_ids), 2)
        short = self.exam.question_ids.filtered(lambda q: q.question_type == 'short_answer')
        mcq = self.exam.question_ids.filtered(lambda q: q.question_type == 'mcq')
        self.assertEqual(short.marks, 5)
        self.assertEqual(short.sample_answer_keywords, 'action, reaction')
        self.assertEqual(len(mcq.option_ids), 3)
        self.assertEqual(mcq.option_ids.filtered('is_correct').option_text, 'Newton')

    def test_xlsx_import_missing_required_column_raises(self):
        content = self._build_xlsx([], headers=['Question Text', 'Marks'])
        wizard = self._wizard_for(content, 'questions.xlsx')
        with self.assertRaises(UserError):
            wizard.action_import()

    def test_xlsx_import_mcq_without_correct_option_raises(self):
        content = self._build_xlsx([
            ['Pick one', 'MCQ', 1, 'A option', 'B option', '', '', '', ''],
        ])
        wizard = self._wizard_for(content, 'questions.xlsx')
        with self.assertRaises(UserError):
            wizard.action_import()

    def test_xlsx_import_unknown_type_raises(self):
        content = self._build_xlsx([
            ['Some question', 'Essay', 5, '', '', '', '', '', ''],
        ])
        wizard = self._wizard_for(content, 'questions.xlsx')
        with self.assertRaises(UserError):
            wizard.action_import()

    def test_unsupported_extension_raises(self):
        wizard = self._wizard_for(b'not a real file', 'questions.txt')
        with self.assertRaises(UserError):
            wizard.action_import()

    # --- Word (.docx) ---

    def test_docx_import_creates_questions(self):
        import docx

        document = docx.Document()
        document.add_paragraph('Q: What is the capital of France?')
        document.add_paragraph('Type: short_answer')
        document.add_paragraph('Marks: 2')
        document.add_paragraph('Answer: Paris; capital of France')
        document.add_paragraph('')
        document.add_paragraph('Q: Who formulated the laws of motion?')
        document.add_paragraph('Type: mcq')
        document.add_paragraph('Marks: 1')
        document.add_paragraph('A) Newton')
        document.add_paragraph('B) Jonas')
        document.add_paragraph('C) Frost')
        document.add_paragraph('Correct: A')
        buffer = io.BytesIO()
        document.save(buffer)

        wizard = self._wizard_for(buffer.getvalue(), 'questions.docx')
        wizard.action_import()

        self.assertEqual(len(self.exam.question_ids), 2)
        mcq = self.exam.question_ids.filtered(lambda q: q.question_type == 'mcq')
        self.assertEqual(len(mcq.option_ids), 3)
        self.assertEqual(mcq.option_ids.filtered('is_correct').option_text, 'Newton')

    # --- shared plain-text block parser (used by both docx and pdf) ---

    def test_parse_text_blocks_short_answer_default_type(self):
        wizard = self.env['bxi.exam.question.import.wizard'].new({'exam_id': self.exam.id})
        rows = wizard._parse_text_blocks('Q: Define photosynthesis\nMarks: 3\nAnswer: sunlight, chlorophyll')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['question_type'], 'short_answer')
        self.assertEqual(rows[0]['marks'], 3.0)
        self.assertEqual(rows[0]['sample_answer_keywords'], 'sunlight, chlorophyll')

    def test_parse_text_blocks_mcq_missing_correct_raises(self):
        wizard = self.env['bxi.exam.question.import.wizard'].new({'exam_id': self.exam.id})
        text = 'Q: Pick one\nType: mcq\nA) Alpha\nB) Beta'
        with self.assertRaises(UserError):
            wizard._parse_text_blocks(text)

    def test_parse_text_blocks_multiple_blocks(self):
        wizard = self.env['bxi.exam.question.import.wizard'].new({'exam_id': self.exam.id})
        text = (
            'Q: First question\nMarks: 1\n\n'
            'Q: Second question\nType: mcq\nMarks: 2\nA) X\nB) Y\nCorrect: B'
        )
        rows = wizard._parse_text_blocks(text)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1]['question_type'], 'mcq')
        self.assertTrue(next(o for o in rows[1]['options'] if o['option_text'] == 'Y')['is_correct'])
